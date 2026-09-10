from __future__ import annotations

import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

import openpyxl


TOOL_DIRECTORY = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_DIRECTORY))

from import_cofid import (  # noqa: E402
    HEADERS,
    WORKSHEET_NAME,
    effective_calories,
    import_workbook,
    parse_nutrient,
    stable_id,
    write_catalogue,
)


class CofidImporterTests(unittest.TestCase):
    def test_trace_is_zero_but_n_is_unavailable(self) -> None:
        self.assertEqual((Decimal(0), True, False), parse_nutrient("Tr"))
        self.assertEqual((None, False, True), parse_nutrient("N"))

    def test_stable_id_is_deterministic_and_independent_of_source_row(self) -> None:
        values = (
            "13-001", "Food", "Description", "A",
            Decimal("84"), Decimal("5.5"), Decimal("12.5"), Decimal("2"),
            "g",
        )
        first = stable_id(*values)
        repeated = stable_id(*values)
        normalized_formatting = stable_id(
            " 13-001 ", " food ", "Description", "a",
            Decimal("84.0"), Decimal("5.50"), Decimal("12.50"), Decimal("2.0"),
            "G",
        )

        self.assertEqual(first, repeated)
        self.assertEqual(first, normalized_formatting)

    def test_stable_id_distinguishes_records_with_duplicate_food_codes(self) -> None:
        first = stable_id(
            "13-669", "Aubergine, roasted", "Vegetable", "A",
            Decimal("50"), Decimal("1"), Decimal("8"), Decimal("2"), "g",
        )
        second = stable_id(
            "13-669", "Watercress, raw", "Vegetable", "A",
            Decimal("20"), Decimal("2"), Decimal("1"), Decimal("0"), "g",
        )

        self.assertNotEqual(first, second)

    def test_imported_ids_do_not_change_when_source_rows_are_reordered(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first_path = root / "first.xlsx"
            reordered_path = root / "reordered.xlsx"
            first_food = [
                "13-001", "First food", "First description", "A",
                "5", "2", "10", "80",
            ]
            second_food = [
                "13-002", "Second food", "Second description", "A",
                "6", "3", "11", "90",
            ]
            self._write_rows(first_path, [first_food, second_food])
            self._write_rows(reordered_path, [second_food, first_food])

            first_import, _ = import_workbook(first_path)
            reordered_import, _ = import_workbook(reordered_path)

            first_ids = {food.name: food.id for food in first_import}
            reordered_ids = {food.name: food.id for food in reordered_import}
            self.assertEqual(first_ids, reordered_ids)

    def test_fractional_calories_use_explicit_half_even_integer_policy(self) -> None:
        self.assertEqual(Decimal("84"), effective_calories(Decimal("84.25")))
        self.assertEqual(Decimal("84"), effective_calories(Decimal("84.5")))
        self.assertEqual(Decimal("86"), effective_calories(Decimal("85.5")))

    def test_fractional_calories_are_mapped_during_import(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "fractional.xlsx"
            self._write_rows(workbook_path, [[
                "13-001", "Fractional food", "Description", "A",
                "5.5", "2.25", "10.75", "84.25",
            ]])

            foods, stats = import_workbook(workbook_path)

            self.assertEqual(1, len(foods))
            food = foods[0]
            self.assertEqual(Decimal("84"), food.calories)
            self.assertEqual(Decimal("5.5"), food.protein)
            self.assertEqual(Decimal("2.25"), food.fat)
            self.assertEqual(Decimal("10.75"), food.carbohydrates)
            self.assertEqual(1, stats.rounded_calories)

    def test_import_validates_and_normalizes_expected_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "fixture.xlsx"
            self._write_fixture(workbook_path)

            foods, stats = import_workbook(workbook_path)

            self.assertEqual(4, stats.source_rows)
            self.assertEqual(3, stats.imported_rows)
            self.assertEqual(1, stats.excluded_rows)
            self.assertEqual(1, stats.trace_conversions)
            self.assertEqual(1, stats.n_values)
            self.assertEqual(1, stats.duplicate_source_codes)
            self.assertEqual(1, stats.duplicate_source_code_rows)
            self.assertEqual(0, stats.rounded_calories)

            by_name = {food.name: food for food in foods}
            self.assertEqual(Decimal("90"), by_name["Fixture food"].calories)
            self.assertEqual(Decimal("5.5"), by_name["Fixture food"].protein)
            self.assertEqual(Decimal(0), by_name["Trace food"].fat)
            self.assertEqual("g", by_name["Fixture food"].serving_unit)
            self.assertEqual("ml", by_name["Red wine"].serving_unit)
            self.assertNotIn("Unavailable food", by_name)
            self.assertEqual(3, len({food.id for food in foods}))

    def test_output_is_byte_for_byte_deterministic(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "fixture.xlsx"
            first_output = root / "first.json"
            second_output = root / "second.json"
            self._write_fixture(workbook_path)
            foods, stats = import_workbook(workbook_path)

            first_hash = write_catalogue(first_output, foods, stats)
            second_hash = write_catalogue(second_output, foods, stats)

            self.assertEqual(first_hash, second_hash)
            self.assertEqual(first_output.read_bytes(), second_output.read_bytes())

    def test_identical_duplicate_records_receive_stable_unique_ids(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "duplicates.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = WORKSHEET_NAME
            headers = list(HEADERS.values())
            sheet.append(headers)
            sheet.append(["metadata"] * len(headers))
            sheet.append(["units"] * len(headers))
            duplicate = [
                "13-001", "Duplicate food", "Same description", "A",
                "5", "2", "10", "80",
            ]
            sheet.append(duplicate)
            sheet.append(duplicate)
            workbook.save(workbook_path)

            foods, _ = import_workbook(workbook_path)

            self.assertEqual(2, len(foods))
            self.assertEqual(2, len({food.id for food in foods}))
            base_id = stable_id(
                "13-001", "Duplicate food", "Same description", "A",
                Decimal("80"), Decimal("5"), Decimal("10"), Decimal("2"),
                "g",
            )
            self.assertEqual([base_id, f"{base_id}-2"], sorted(
                (food.id for food in foods),
                key=len,
            ))

    def test_missing_header_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "invalid.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = WORKSHEET_NAME
            sheet.append(["Food Code"])
            workbook.save(workbook_path)

            with self.assertRaisesRegex(ValueError, "Missing expected header"):
                import_workbook(workbook_path)

    def test_missing_expected_worksheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "invalid.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Unexpected"
            workbook.save(workbook_path)

            with self.assertRaisesRegex(ValueError, "Missing worksheet"):
                import_workbook(workbook_path)

    @staticmethod
    def _write_fixture(path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = WORKSHEET_NAME
        headers = list(HEADERS.values())
        sheet.append(headers)
        sheet.append(["metadata"] * len(headers))
        sheet.append(["units"] * len(headers))
        sheet.append([
            "13-001", "Fixture food", "Fixture description", "A",
            "5.5", "2.0", "12.5", "90",
        ])
        sheet.append([
            "13-002", "Trace food", "Fixture description", "A",
            "1", "Tr", "2", "20",
        ])
        sheet.append([
            "13-003", "Unavailable food", "Fixture description", "A",
            "1", "2", "N", "20",
        ])
        sheet.append([
            "13-001", "Red wine", "Fixture description", "QE",
            "0.1", "0", "2.5", "75",
        ])
        workbook.save(path)

    @staticmethod
    def _write_rows(path: Path, rows: list[list[str]]) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = WORKSHEET_NAME
        headers = list(HEADERS.values())
        sheet.append(headers)
        sheet.append(["metadata"] * len(headers))
        sheet.append(["units"] * len(headers))

        for row in rows:
            sheet.append(row)

        workbook.save(path)

if __name__ == "__main__":
    unittest.main()
