"""Build the normalized Comfy Capy CoFID 2021 catalogue artifact."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from pathlib import Path
from typing import Any

import openpyxl


WORKSHEET_NAME = "1.3 Proximates"
FIRST_DATA_ROW = 4
SOURCE_NAME = "CoFID"
SOURCE_VERSION = "2021"
SOURCE_URL = (
    "https://www.gov.uk/government/publications/"
    "composition-of-foods-integrated-dataset-cofid"
)
ALCOHOLIC_BEVERAGE_GROUPS = frozenset({
    "Q", "QA", "QC", "QE", "QF", "QG", "QI", "QK"
})

HEADERS = {
    "food_code": "Food Code",
    "name": "Food Name",
    "description": "Description",
    "group": "Group",
    "protein": "Protein (g)",
    "fat": "Fat (g)",
    "carbohydrates": "Carbohydrate (g)",
    "calories": "Energy (kcal) (kcal)",
}


@dataclass(frozen=True)
class ImportStats:
    source_rows: int
    imported_rows: int
    excluded_rows: int
    trace_conversions: int
    n_values: int
    duplicate_source_codes: int
    duplicate_source_code_rows: int
    rounded_calories: int


@dataclass(frozen=True)
class FoodRecord:
    id: str
    source_code: str
    source_row: int
    name: str
    description: str
    group: str
    calories: Decimal
    protein: Decimal
    carbohydrates: Decimal
    fat: Decimal
    serving_size: int
    serving_unit: str


def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(str(value).split())


def parse_nutrient(value: Any) -> tuple[Decimal | None, bool, bool]:
    """Return value, whether Tr was converted, and whether N was encountered."""
    text = clean_text(value)
    marker = text.casefold()

    if marker == "tr":
        return Decimal(0), True, False

    if marker == "n":
        return None, False, True

    if not text:
        return None, False, False

    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, False, False

    if not number.is_finite() or number < 0:
        return None, False, False

    return number, False, False


def canonical_decimal(value: Decimal) -> str:
    if value == 0:
        return "0"

    normalized = value.normalize()
    return format(normalized, "f")


def effective_calories(value: Decimal) -> Decimal:
    """Map source kcal to the integer domain model using .NET's default rule."""
    return value.quantize(Decimal("1"), rounding=ROUND_HALF_EVEN)


def stable_id(
    source_code: str,
    name: str,
    description: str,
    group: str,
    calories: Decimal,
    protein: Decimal,
    carbohydrates: Decimal,
    fat: Decimal,
    serving_unit: str,
) -> str:
    # Row position is provenance only. Identity comes from normalized immutable
    # record content so source reordering does not change catalogue IDs.
    identity = json.dumps(
        [
            clean_text(source_code).casefold(),
            clean_text(name).casefold(),
            clean_text(description).casefold(),
            clean_text(group).upper(),
            canonical_decimal(calories),
            canonical_decimal(protein),
            canonical_decimal(carbohydrates),
            canonical_decimal(fat),
            clean_text(serving_unit).casefold(),
        ],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    digest = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:20]
    return f"cf21-{digest}"


def import_workbook(path: Path) -> tuple[list[FoodRecord], ImportStats]:
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    if WORKSHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Missing worksheet: {WORKSHEET_NAME}")

    sheet = workbook[WORKSHEET_NAME]
    header_values = [
        clean_text(cell.value)
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    ]

    duplicate_headers = [
        header
        for header, count in Counter(header_values).items()
        if header and count > 1
    ]
    if duplicate_headers:
        joined = ", ".join(sorted(duplicate_headers))
        workbook.close()
        raise ValueError(f"Duplicate worksheet headers: {joined}")

    indexes: dict[str, int] = {}
    for field, header in HEADERS.items():
        if header not in header_values:
            workbook.close()
            raise ValueError(f"Missing expected header: {header}")
        indexes[field] = header_values.index(header)

    foods: list[FoodRecord] = []
    source_codes: list[str] = []
    seen_ids: set[str] = set()
    source_rows = 0
    excluded_rows = 0
    trace_conversions = 0
    n_values = 0
    rounded_calories = 0
    identity_counts: Counter[str] = Counter()

    for source_row, row in enumerate(
        sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True),
        FIRST_DATA_ROW,
    ):
        selected = {
            field: row[index]
            for field, index in indexes.items()
        }

        if not any(clean_text(value) for value in selected.values()):
            continue

        source_rows += 1
        source_code = clean_text(selected["food_code"])
        name = clean_text(selected["name"])
        description = clean_text(selected["description"])
        group = clean_text(selected["group"]).upper()
        source_codes.append(source_code)

        if group.startswith("Q") and group not in ALCOHOLIC_BEVERAGE_GROUPS:
            workbook.close()
            raise ValueError(
                f"Unexpected Q-family group '{group}' at row {source_row}"
            )

        nutrients: dict[str, Decimal] = {}
        reliable = bool(source_code and name and group)

        for field in ("protein", "fat", "carbohydrates", "calories"):
            nutrient, was_trace, was_n = parse_nutrient(selected[field])
            trace_conversions += int(was_trace)
            n_values += int(was_n)

            if nutrient is None:
                reliable = False
            else:
                nutrients[field] = nutrient

        if not reliable:
            excluded_rows += 1
            continue

        unit = (
            "ml"
            if group in ALCOHOLIC_BEVERAGE_GROUPS
            else "g"
        )
        calories = effective_calories(nutrients["calories"])
        rounded_calories += int(calories != nutrients["calories"])
        base_identifier = stable_id(
            source_code,
            name,
            description,
            group,
            calories,
            nutrients["protein"],
            nutrients["carbohydrates"],
            nutrients["fat"],
            unit,
        )
        identity_counts[base_identifier] += 1
        occurrence = identity_counts[base_identifier]
        identifier = (
            base_identifier
            if occurrence == 1
            else f"{base_identifier}-{occurrence}"
        )

        if identifier in seen_ids:
            workbook.close()
            raise ValueError(f"Generated duplicate catalogue ID: {identifier}")
        seen_ids.add(identifier)

        foods.append(FoodRecord(
            id=identifier,
            source_code=source_code,
            source_row=source_row,
            name=name,
            description=description,
            group=group,
            calories=calories,
            protein=nutrients["protein"],
            carbohydrates=nutrients["carbohydrates"],
            fat=nutrients["fat"],
            serving_size=100,
            serving_unit=unit,
        ))

    foods.sort(key=lambda food: food.id)
    duplicate_counts = [
        count
        for count in Counter(source_codes).values()
        if count > 1
    ]
    stats = ImportStats(
        source_rows=source_rows,
        imported_rows=len(foods),
        excluded_rows=excluded_rows,
        trace_conversions=trace_conversions,
        n_values=n_values,
        duplicate_source_codes=len(duplicate_counts),
        duplicate_source_code_rows=sum(count - 1 for count in duplicate_counts),
        rounded_calories=rounded_calories,
    )
    workbook.close()
    return foods, stats


def json_compatible(record: FoodRecord) -> dict[str, Any]:
    values = asdict(record)

    for key in ("calories", "protein", "carbohydrates", "fat"):
        number = values[key]
        values[key] = int(number) if number == number.to_integral() else float(number)

    return values


def write_catalogue(
    output_path: Path,
    foods: list[FoodRecord],
    stats: ImportStats,
) -> str:
    document = {
        "provider": "cofid",
        "source": SOURCE_NAME,
        "version": SOURCE_VERSION,
        "sourceUrl": SOURCE_URL,
        "import": {
            "worksheet": WORKSHEET_NAME,
            "firstDataRow": FIRST_DATA_ROW,
            "sourceRows": stats.source_rows,
            "importedRows": stats.imported_rows,
            "excludedRows": stats.excluded_rows,
            "traceConversions": stats.trace_conversions,
            "nValues": stats.n_values,
            "duplicateSourceCodes": stats.duplicate_source_codes,
            "duplicateSourceCodeRows": stats.duplicate_source_code_rows,
            "roundedCalories": stats.rounded_calories,
        },
        "foods": [json_compatible(food) for food in foods],
    }
    serialized = json.dumps(
        document,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ) + "\n"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(output_path.suffix + ".tmp")
    temporary_path.write_text(serialized, encoding="utf-8", newline="\n")
    temporary_path.replace(output_path)
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def main() -> int:
    repository_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=repository_root / "tools/cofid/source/CoFID_2021.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=(
            repository_root
            / "CalorieTracker/Data/Catalogues/cofid-2021.json"
        ),
    )
    arguments = parser.parse_args()

    foods, stats = import_workbook(arguments.source)
    digest = write_catalogue(arguments.output, foods, stats)

    print(f"Source rows: {stats.source_rows}")
    print(f"Imported records: {stats.imported_rows}")
    print(f"Excluded incomplete/unreliable records: {stats.excluded_rows}")
    print(f"Core Tr values converted to zero: {stats.trace_conversions}")
    print(f"Core N values left unavailable: {stats.n_values}")
    print(f"Fractional kcal values rounded to integers: {stats.rounded_calories}")
    print(
        "Duplicate source Food Codes: "
        f"{stats.duplicate_source_codes} codes, "
        f"{stats.duplicate_source_code_rows} additional rows"
    )
    print(f"Artifact SHA-256: {digest}")
    print(f"Wrote: {arguments.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
